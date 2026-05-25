#!/usr/bin/env python3
"""
Build a DOCX from `ETC_PAPER_DRAFT.md`.

Goal: provide a portable artifact you can open on any device for review.
This is a light Markdown-to-DOCX converter (headings + paragraphs + basic bullets).
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Pt


ROOT = Path(__file__).resolve().parent
SRC = ROOT / "ETC_PAPER_DRAFT.md"
OUT = ROOT / "ETC_PAPER_DRAFT.docx"


HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
BULLET_RE = re.compile(r"^(\s*)([-*])\s+(.*)$")


def set_normal_style(doc: Document) -> None:
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Times New Roman"
    font.size = Pt(12)


def add_paragraph_md(doc: Document, text: str) -> None:
    # Keep it simple: strip Markdown emphasis markers but preserve content.
    text = text.replace("**", "").replace("__", "").replace("`", "")
    doc.add_paragraph(text)


def main() -> None:
    if not SRC.exists():
        raise SystemExit(f"Missing source: {SRC}")

    doc = Document()
    set_normal_style(doc)

    lines = SRC.read_text(encoding="utf-8").splitlines()
    buf: list[str] = []

    def flush_paragraph() -> None:
        nonlocal buf
        if buf:
            add_paragraph_md(doc, " ".join(s.strip() for s in buf).strip())
            buf = []

    for raw in lines:
        line = raw.rstrip()

        if not line.strip():
            flush_paragraph()
            continue

        heading = HEADING_RE.match(line)
        if heading:
            flush_paragraph()
            level = len(heading.group(1))
            text = heading.group(2).strip().replace("**", "")
            # Map to Word heading levels 1..3, then fall back.
            if level <= 3:
                doc.add_heading(text, level=level)
            else:
                # For deeper headings, use bold paragraph.
                p = doc.add_paragraph()
                run = p.add_run(text)
                run.bold = True
            continue

        bullet = BULLET_RE.match(line)
        if bullet:
            flush_paragraph()
            indent = bullet.group(1) or ""
            text = bullet.group(3).strip().replace("**", "")
            style = "List Bullet" if len(indent) < 2 else "List Bullet 2"
            doc.add_paragraph(text, style=style)
            continue

        buf.append(line)

    flush_paragraph()

    doc.save(OUT)
    print(str(OUT))


if __name__ == "__main__":
    main()

